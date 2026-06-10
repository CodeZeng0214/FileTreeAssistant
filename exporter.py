"""导出模块 —— 支持 TXT 文本导出和 Excel 表格导出"""

import os
from tkinter import filedialog, messagebox
from scanner import FileNode, format_as_tree, flatten_nodes


def export_txt(nodes: list[FileNode], root_path: str, parent_window=None) -> bool:
    """
    将目录树导出为 TXT 文件。

    返回: 是否导出成功
    """
    tree_text = format_as_tree(nodes, root_path)
    
    file_path = filedialog.asksaveasfilename(
        title="导出 TXT 文件",
        defaultextension=".txt",
        filetypes=[("文本文件", "*.txt"), ("所有文件", "*.*")],
        initialfile=f"{os.path.basename(root_path.rstrip(os.sep))}_目录树.txt",
        parent=parent_window,
    )
    if not file_path:
        return False

    try:
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(tree_text)
        return True
    except Exception as e:
        messagebox.showerror("导出失败", f"导出 TXT 时出错：\n{e}", parent=parent_window)
        return False


def export_excel(
    nodes: list[FileNode], 
    root_path: str, 
    full_info: bool = False,
    parent_window=None,
) -> bool:
    """
    将目录树导出为 Excel 文件。

    参数:
        nodes:       节点列表
        root_path:   根路径
        full_info:   True=完整信息（含大小、日期），False=基础信息
        parent_window: 父窗口

    返回: 是否导出成功
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    file_path = filedialog.asksaveasfilename(
        title="导出 Excel 文件",
        defaultextension=".xlsx",
        filetypes=[("Excel 文件", "*.xlsx"), ("所有文件", "*.*")],
        initialfile=f"{os.path.basename(root_path.rstrip(os.sep))}_目录树.xlsx",
        parent=parent_window,
    )
    if not file_path:
        return False

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "目录树"

        # 表头样式
        header_fill = PatternFill(start_color="BBDEFB", end_color="BBDEFB", fill_type="solid")  # 淡蓝
        header_font = Font(name="微软雅黑", bold=True, size=11, color="1565C0")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="B0BEC5"),
            right=Side(style="thin", color="B0BEC5"),
            top=Side(style="thin", color="B0BEC5"),
            bottom=Side(style="thin", color="B0BEC5"),
        )
        cell_font = Font(name="微软雅黑", size=10)
        cell_alignment = Alignment(vertical="center")
        level_alignment = Alignment(horizontal="center", vertical="center")

        if full_info:
            headers = ["层级", "名称", "类型", "大小", "修改日期", "完整路径"]
        else:
            headers = ["层级", "名称", "类型", "完整路径"]

        # 写表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # 写数据
        flat_nodes = flatten_nodes(nodes)
        for row_idx, node in enumerate(flat_nodes, 2):
            node_type = "📁 文件夹" if node.is_dir else "📄 文件"
            
            if full_info:
                size_str = _format_size(node.size) if not node.is_dir else ""
                row_data = [
                    node.level,
                    node.name,
                    node_type,
                    size_str,
                    node.modified_time,
                    node.path,
                ]
            else:
                row_data = [
                    node.level,
                    node.name,
                    node_type,
                    node.path,
                ]

            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = cell_font
                cell.border = thin_border
                if col_idx == 1:
                    cell.alignment = level_alignment
                else:
                    cell.alignment = cell_alignment

        # 自动调整列宽
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    # 中文字符按2个宽度计算
                    cell_len = sum(2 if ord(c) > 127 else 1 for c in str(cell.value or ""))
                    max_length = max(max_length, cell_len)
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 4, 60)

        # 冻结首行
        ws.freeze_panes = "A2"

        wb.save(file_path)
        return True

    except Exception as e:
        messagebox.showerror("导出失败", f"导出 Excel 时出错：\n{e}", parent=parent_window)
        return False


def _format_size(size_bytes: int) -> str:
    """字节转人类可读大小"""
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    elif size_bytes < 1024 * 1024 * 1024:
        return f"{size_bytes / (1024 * 1024):.1f} MB"
    else:
        return f"{size_bytes / (1024 * 1024 * 1024):.2f} GB"
